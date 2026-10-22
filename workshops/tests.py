import shutil
import tempfile
from datetime import timedelta
from pathlib import Path
from unittest import mock

from django.contrib.auth import get_user_model
from django.core import mail, signing
from django.core.files.base import ContentFile
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from . import storage
from .models import Announcement, Feedback, GalleryImage, Unsubscribed, Workshop

TMP = Path(tempfile.mkdtemp())


def png(name="pay.png"):
    import io
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (40, 40), "green").save(buf, "PNG")
    return SimpleUploadedFile(name, buf.getvalue(), content_type="image/png")


def cert_template_bytes():
    import io
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (1600, 1130), "#fffdf5").save(buf, "PNG")
    return buf.getvalue()


def login_staff(client):
    get_user_model().objects.create_user("s", password="pw12345678", is_staff=True, email="staff@example.com")
    client.login(username="s", password="pw12345678")


@override_settings(
    REGISTRATIONS_XLSX=TMP / "registrations.xlsx", PAYMENTS_DIR=TMP / "payments", CERT_TEMPLATES_DIR=TMP / "certs", ANNOUNCE_FILES_DIR=TMP / "announce",
    MEDIA_ROOT=TMP / "media",
    EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend",
    EMAIL_HOST_USER="sender@example.com", EMAIL_HOST_PASSWORD="x", ADMIN_NOTIFY_EMAIL="admin@example.com",
    ALLOWED_HOSTS=["testserver"], SITE_URL="https://workshops.example.com",
)
class WorkshopFlowTests(TestCase):
    def setUp(self):
        cache.clear()
        p = storage.path()
        for f in (p, Path(str(p) + ".lock")):
            if f.exists():
                f.unlink()
        now = timezone.now()
        self.w = Workshop.objects.create(
            title="AI Prompt Engineering", description="Learn prompts", start_datetime=now + timedelta(days=7),
            end_datetime=now + timedelta(days=7, hours=2), mode="online", meeting_link="https://meet.example/abc",
            status="published", capacity=1, fee=9, payment_upi="x@upi", featured=True,
        )

    def post(self, **over):
        data = {
            "workshop": str(self.w.pk), "full_name": "Test User", "email": "Test@Example.com", "phone": "98765 43210",
            "college": "ABC College", "payment_reference": "UTR123",
            "current_address": "12 Main Road, Dindigul", "education_qualification": "B.E. CSE", "current_status": "3rd year student",
            "heard_from": "Instagram", "interest_domain": "AI and Data", "skill_level": "Beginner",
            "demo_interest": "Maybe", "demo_mode": "Online", "demo_slot": "Evening", "queries": "", "referred_by": "A friend",
            "payment_screenshot": png(), "wants_certificate": "yes",
            "consent": "on", "website": "", "ts": signing.dumps(__import__("time").time() - 10, salt="reg-form"),
        }
        data.update(over)
        return self.client.post(reverse("register"), data)

    def rows(self):
        return load_workbook(storage.path())["Registrations"]

    # --- public pages
    def test_public_pages_render(self):
        for name in ("home", "workshop_list", "register", "gallery"):
            self.assertEqual(self.client.get(reverse(name)).status_code, 200, name)
        self.assertEqual(self.client.get(self.w.get_absolute_url()).status_code, 200)
        self.assertContains(self.client.get(reverse("home")), "Errors2Experts Weekly Workshops")

    def test_meeting_link_never_public(self):
        for url in (reverse("home"), self.w.get_absolute_url(), reverse("register")):
            self.assertNotContains(self.client.get(url), "meet.example")

    def test_draft_hidden_but_staff_can_preview(self):
        self.w.status = "draft"; self.w.save()
        self.assertEqual(self.client.get(self.w.get_absolute_url()).status_code, 404)
        get_user_model().objects.create_user("s", password="pw12345678", is_staff=True)
        self.client.login(username="s", password="pw12345678")
        self.assertContains(self.client.get(self.w.get_absolute_url()), "Draft preview")

    def test_remaining_seats_only_when_capacity_set(self):
        self.assertContains(self.client.get(reverse("home")), "1 seat left")
        self.w.capacity = None; self.w.save()
        self.assertNotRegex(self.client.get(reverse("home")).content.decode(), r"seats? left")

    # --- registration
    def test_success_saves_to_excel_and_emails(self):
        r = self.post()
        self.assertRedirects(r, reverse("thank_you"))
        ws = self.rows()
        self.assertEqual(ws.max_row, 2)
        vals = {c.value: ws.cell(row=2, column=i + 1).value for i, c in enumerate(ws[1])}
        self.assertEqual(vals["email"], "test@example.com")
        self.assertEqual(vals["phone"], "9876543210")
        self.assertEqual(vals["email_status"], "sent")
        self.assertEqual(len(mail.outbox), 2)
        self.assertEqual(mail.outbox[0].subject, "Registration Confirmed – AI Prompt Engineering")
        self.assertEqual(mail.outbox[1].subject, "New Workshop Registration – AI Prompt Engineering")
        self.assertIn("meet.example/abc", mail.outbox[0].body)  # joining link goes only to the registrant
        ty = self.client.get(reverse("thank_you"))
        self.assertContains(ty, "You're registered")
        self.assertContains(ty, "https://errors2experts.com/")

    def test_nothing_in_database_for_registrations(self):
        self.post()
        from django.apps import apps
        names = {m.__name__ for m in apps.get_app_config("workshops").get_models()}
        self.assertEqual(names, {"Workshop", "GalleryImage", "Feedback", "Unsubscribed", "Announcement"})  # no registration model

    def test_field_errors(self):
        r = self.post(email="bad", phone="123", consent="")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Enter a valid email address.")
        self.assertContains(r, "Enter a valid 10-digit mobile number.")
        self.assertContains(r, "Consent is required")
        self.assertFalse(storage.path().exists() and self.rows().max_row > 1)

    def test_payment_screenshot_required_when_fee(self):
        r = self.post(payment_screenshot="")
        self.assertContains(r, "Attach your payment screenshot")
        self.assertFalse(storage.path().exists() and self.rows().max_row > 1)

    def test_free_workshop_needs_no_screenshot(self):
        self.w.fee = None; self.w.save()
        self.assertRedirects(self.post(payment_screenshot=""), reverse("thank_you"))

    def test_bad_screenshot_rejected(self):
        bad = SimpleUploadedFile("x.png", b"not an image", content_type="image/png")
        self.assertContains(self.post(payment_screenshot=bad), "Upload a screenshot")

    def test_screenshot_saved_privately_and_visible_only_to_staff(self):
        self.post()
        ws = self.rows()
        reg_id = ws.cell(row=2, column=1).value
        name = ws.cell(row=2, column=17).value
        self.assertEqual(name, f"{reg_id}.png")
        self.assertTrue((TMP / "payments" / name).is_file())
        url = reverse("payment_proof", args=[reg_id])
        self.assertEqual(self.client.get(url).status_code, 302)  # anonymous: sent to login
        get_user_model().objects.create_user("s", password="pw12345678", is_staff=True)
        self.client.login(username="s", password="pw12345678")
        self.assertEqual(self.client.get(url).status_code, 200)
        self.assertEqual(self.client.get(reverse("payment_proof", args=["E2E-20200101-AAAAAA"])).status_code, 404)
        self.assertContains(self.client.get(reverse("registrations")), "Screenshot")

    def test_screenshot_removed_when_registration_rejected(self):
        self.post()
        before = set((TMP / "payments").iterdir())
        self.post(email="b@example.com")  # workshop is full
        self.assertEqual(set((TMP / "payments").iterdir()), before)

    def test_payment_qr_from_upi_when_no_image(self):
        self.assertIn('"payqr": "/pay-qr/', self.client.get(reverse("register")).content.decode().replace("&quot;", '"'))
        r = self.client.get(reverse("pay_qr", args=[self.w.pk]))
        self.assertTrue(r.content.startswith(b"\x89PNG"))
        self.w.fee = None; self.w.save()
        self.assertEqual(self.client.get(reverse("pay_qr", args=[self.w.pk])).status_code, 404)

    # --- feedback and gallery
    def fb(self, **over):
        data = {"workshop": str(self.w.pk), "name": "Meena", "email": "Meena@College.edu", "rating": "5",
                "message": "Very practical session, loved it.", "website": "",
                "ts": signing.dumps(__import__("time").time() - 10, salt="reg-form")}
        data.update(over)
        return self.client.post(reverse("feedback"), data)

    def test_feedback_flow_and_public_display(self):
        self.assertEqual(self.client.get(reverse("feedback")).status_code, 200)
        self.assertRedirects(self.fb(), reverse("feedback_thanks"))
        f = Feedback.objects.get()
        self.assertEqual((f.email, f.rating, f.workshop_title), ("meena@college.edu", 5, "AI Prompt Engineering"))
        home = self.client.get(reverse("home")).content.decode()
        self.assertIn("Very practical session", home)
        self.assertIn("Meena", home)
        self.assertNotIn("meena@college.edu", home)  # email never public

    def test_feedback_validation_and_duplicates(self):
        self.assertContains(self.fb(rating=""), "Choose a rating")
        self.assertContains(self.fb(rating="9"), "Select a valid choice")
        self.assertContains(self.fb(email="nope"), "valid email")
        self.assertContains(self.fb(message="short"), "at least")
        self.assertContains(self.fb(website="spam"), "rejected")
        self.assertEqual(Feedback.objects.count(), 0)
        self.fb()
        self.assertContains(self.fb(), "already submitted")
        self.assertEqual(Feedback.objects.count(), 1)

    def test_admin_can_delete_feedback_on_website(self):
        self.fb()
        pk = Feedback.objects.get().pk
        self.assertEqual(self.client.post(reverse("feedback_delete", args=[pk])).status_code, 302)
        self.assertEqual(Feedback.objects.count(), 1)  # anonymous cannot delete
        self.assertEqual(self.client.get(reverse("feedback_admin")).status_code, 302)
        get_user_model().objects.create_user("s", password="pw12345678", is_staff=True)
        self.client.login(username="s", password="pw12345678")
        self.assertContains(self.client.get(reverse("feedback_admin")), "meena@college.edu")
        self.assertEqual(self.client.get(reverse("feedback_delete", args=[pk])).status_code, 405)  # POST only
        self.client.post(reverse("feedback_delete", args=[pk]))
        self.assertEqual(Feedback.objects.count(), 0)
        self.assertNotContains(self.client.get(reverse("home")), "Very practical")

    def test_hidden_feedback_not_shown(self):
        self.fb()
        Feedback.objects.update(show_on_site=False)
        self.assertNotContains(self.client.get(reverse("home")), "Very practical")

    def test_gallery_on_home(self):
        GalleryImage.objects.create(kind="attendees", image=png("a.png"), caption="Batch photo", workshop=self.w)
        GalleryImage.objects.create(kind="feedback", image=png("b.png"), caption="WhatsApp feedback")
        home = self.client.get(reverse("home")).content.decode()
        self.assertIn("Batch photo", home)
        self.assertIn("WhatsApp feedback", home)  # feedback screenshot caption/alt shows in "What students say"
        self.assertIn("What students say", home)
        self.assertIn("Our previous workshops", home)
        self.assertContains(self.client.get(self.w.get_absolute_url()), "Photos from this workshop")

    def test_duplicate_blocked(self):
        self.w.capacity = None; self.w.save()
        self.post()
        self.assertContains(self.post(email="TEST@example.com"), "already registered")
        self.assertEqual(self.rows().max_row, 2)

    def test_capacity_enforced(self):
        self.post()
        r = self.post(email="b@example.com")
        self.assertContains(r, "full or closed")
        self.assertEqual(self.rows().max_row, 2)
        listing = self.client.get(reverse("workshop_list"))
        self.assertContains(listing, "pill-full")
        self.assertNotContains(listing, "register/?w=")  # no per-workshop Register button when full

    def test_capacity_enforced_inside_the_lock_even_if_page_state_is_stale(self):
        row = {"registration_id": "A", "workshop_id": self.w.pk, "email": "a@x.com", "registration_status": "confirmed"}
        storage.add_registration(row, capacity=1)
        with self.assertRaises(storage.FullError):
            storage.add_registration({**row, "registration_id": "B", "email": "b@x.com"}, capacity=1)
        with self.assertRaises(storage.DuplicateError):
            storage.add_registration({**row, "registration_id": "C"}, capacity=5)
        self.assertEqual(self.rows().max_row, 2)

    def test_closed_and_expired(self):
        self.w.registration_status = "closed"; self.w.save()
        self.assertContains(self.client.get(reverse("register")), "No workshops are open")
        self.w.registration_status = "open"
        self.w.start_datetime = timezone.now() - timedelta(days=3)
        self.w.end_datetime = timezone.now() - timedelta(days=3) + timedelta(hours=2)
        self.w.save()
        self.assertContains(self.client.get(reverse("register")), "No workshops are open")
        self.assertContains(self.client.get(reverse("workshop_list")), "Completed")

    def test_spam_protection(self):
        self.assertContains(self.post(website="http://spam"), "Submission rejected")
        fast = signing.dumps(__import__("time").time(), salt="reg-form")
        self.assertContains(self.post(ts=fast), "very fast")
        self.assertContains(self.post(ts="forged"), "expired")
        self.assertFalse(storage.path().exists() and self.rows().max_row > 1)

    def test_multiple_open_workshops_not_preselected(self):
        Workshop.objects.create(title="Second", start_datetime=timezone.now() + timedelta(days=9), status="published")
        html = self.client.get(reverse("register")).content.decode()
        self.assertIn("Select a workshop…", html)
        options = html.split('id="id_workshop"')[1].split("</select>")[0]
        self.assertIn('<option value="" selected>', options)  # nothing chosen for the student
        self.assertNotIn('value="1" selected', options)
        self.assertNotIn('value="2" selected', options)

    def test_hybrid_requires_mode(self):
        self.w.mode = "hybrid"; self.w.venue = "Hall"; self.w.capacity = None; self.w.save()
        self.assertContains(self.post(), "Choose online or offline")
        self.assertRedirects(self.post(participation_mode="offline"), reverse("thank_you"))

    def test_email_failure_keeps_registration_and_retry_does_not_duplicate(self):
        with mock.patch("workshops.mailer._send", side_effect=RuntimeError("smtp down")):
            self.assertRedirects(self.post(), reverse("thank_you"))
        ws = self.rows()
        self.assertTrue(str(ws.cell(row=2, column=14).value).startswith("failed"))
        self.assertContains(self.client.get(reverse("thank_you")), "could not send the confirmation email")
        get_user_model().objects.create_user("s", password="pw12345678", is_staff=True)
        self.client.login(username="s", password="pw12345678")
        reg_id = ws.cell(row=2, column=1).value
        self.client.post(reverse("retry_email", args=[reg_id]))
        self.assertEqual(str(self.rows().cell(row=2, column=14).value), "sent")
        n = len(mail.outbox)
        self.client.post(reverse("retry_email", args=[reg_id]))
        self.assertEqual(len(mail.outbox), n)  # no second confirmation

    def test_excel_locked_by_user_is_reported_honestly(self):
        self.post(email="first@example.com")
        self.w.capacity = None; self.w.save()
        with mock.patch("workshops.storage.os.replace", side_effect=PermissionError):
            r = self.post(email="second@example.com")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "open in Excel")
        self.assertEqual(self.rows().max_row, 2)  # second was NOT saved and NOT reported as success
        self.assertFalse(Path(str(storage.path()).replace(".xlsx", ".saving.xlsx")).exists())

    def test_formula_injection_stored_as_text(self):
        self.w.capacity = None; self.w.save()
        self.post(full_name="=HYPERLINK(\"http://evil\")", college="+SUM(1)")
        c = self.rows().cell(row=2, column=5)
        self.assertEqual(c.data_type, "s")
        self.assertTrue(c.value.startswith("="))

    # --- staff areas
    def test_staff_pages_require_login(self):
        for name in ("dashboard", "registrations", "export_excel", "export_csv", "qr_admin", "feedback_admin",
                     "certificates", "announce_index"):
            r = self.client.get(reverse(name))
            self.assertEqual(r.status_code, 302, name)
            self.assertIn("/admin/login/", r["Location"])
        self.assertEqual(self.client.post(reverse("retry_email", args=["x"])).status_code, 302)

    def test_dashboard_filters_exports(self):
        self.w.capacity = None; self.w.save()
        self.post()
        self.post(email="b@example.com", full_name="Bala", college="XYZ Tech")
        get_user_model().objects.create_user("s", password="pw12345678", is_staff=True)
        self.client.login(username="s", password="pw12345678")
        self.assertContains(self.client.get(reverse("dashboard")), ">2<")
        r = self.client.get(reverse("registrations") + "?q=bala")
        self.assertContains(r, "Bala"); self.assertNotContains(r, "Test User")
        r = self.client.get(reverse("registrations") + "?college=abc")
        self.assertContains(r, "Test User"); self.assertNotContains(r, "Bala")
        x = self.client.get(reverse("export_excel") + "?q=bala")
        self.assertIn("spreadsheetml", x["Content-Type"])
        import io
        wb = load_workbook(io.BytesIO(x.content))["Registrations"]
        self.assertEqual(wb.max_row, 2)
        self.assertIn("registration_id", self.client.get(reverse("export_csv")).content.decode())

    # --- QR
    def test_qr_codes(self):
        r = self.client.get(reverse("qr_png", args=["permanent"]))
        self.assertEqual(r["Content-Type"], "image/png")
        self.assertTrue(r.content.startswith(b"\x89PNG"))
        self.assertIn("attachment", self.client.get(reverse("qr_png", args=["permanent"]) + "?download=1")["Content-Disposition"])
        self.assertEqual(self.client.get(reverse("qr_png", args=[self.w.slug])).status_code, 200)
        self.assertEqual(self.client.get(reverse("qr_print", args=["permanent"])).status_code, 200)
        self.w.status = "draft"; self.w.save()
        self.assertEqual(self.client.get(reverse("qr_png", args=[self.w.slug])).status_code, 404)

    def test_permanent_qr_url_is_stable(self):
        from .views import _qr_target
        before = _qr_target("permanent")[0]
        Workshop.objects.create(title="Next week", start_datetime=timezone.now() + timedelta(days=14), status="published", featured=True)
        self.assertEqual(_qr_target("permanent")[0], before)
        self.assertEqual(before, "https://workshops.example.com/register/")

    # --- Google-form fields
    def test_google_form_fields_are_required_and_saved(self):
        r = self.post(heard_from="x", skill_level="", referred_by="")
        self.assertContains(r, "Please fill this in.")
        self.assertContains(r, "This field is required.")
        self.assertFalse(storage.path().exists() and self.rows().max_row > 1)
        self.assertRedirects(self.post(), reverse("thank_you"))
        ws = self.rows()
        vals = {c.value: ws.cell(row=2, column=i + 1).value for i, c in enumerate(ws[1])}
        self.assertEqual(vals["skill_level"], "Beginner")
        self.assertEqual(vals["referred_by"], "A friend")
        self.assertEqual(vals["education_qualification"], "B.E. CSE")
        self.assertEqual(vals["payment_status"], "pending")

    # --- payment tick, feedback and certificates
    def add_template(self, w=None):
        w = w or self.w
        w.cert_template.save("template.png", ContentFile(cert_template_bytes()), save=True)
        return w

    def start_workshop(self):
        """Registration closes when the workshop starts, so register first, then move the start into the past."""
        now = timezone.now()
        Workshop.objects.filter(pk=self.w.pk).update(start_datetime=now - timedelta(hours=3), end_datetime=now - timedelta(hours=1))

    def reg_id(self):
        return self.rows().cell(row=2, column=1).value

    def cert_status(self):
        ws = self.rows()
        headers = [c.value for c in ws[1]]
        return ws.cell(row=2, column=headers.index("certificate_status") + 1).value

    def certificate_mails(self):
        return [m for m in mail.outbox if m.subject.startswith("Your certificate")]

    def test_certificate_sent_after_payment_tick_then_feedback(self):
        self.add_template(); self.post(); self.start_workshop(); login_staff(self.client)
        rid = self.reg_id()
        self.client.post(reverse("payment_set", args=[rid]), {"verified": "on"})
        self.assertEqual(self.certificate_mails(), [])  # payment ok, but no feedback yet
        self.assertEqual(self.client.get(reverse("registrations")).status_code, 200)
        self.client.logout()
        self.assertRedirects(self.fb(email="Test@Example.com"), reverse("feedback_thanks"))
        sent = self.certificate_mails()
        self.assertEqual(len(sent), 1)
        self.assertEqual(sent[0].to, ["test@example.com"])
        name, data, mime = sent[0].attachments[0]
        self.assertTrue(name.endswith(".pdf")); self.assertTrue(data.startswith(b"%PDF")); self.assertEqual(mime, "application/pdf")
        self.assertEqual(self.cert_status(), "sent")
        self.assertContains(self.client.get(reverse("feedback_thanks")), "has been emailed")

    def test_certificate_sent_when_payment_ticked_after_feedback(self):
        self.add_template(); self.post(); self.start_workshop()
        self.fb(email="Test@Example.com")
        self.assertEqual(self.certificate_mails(), [])  # feedback in, payment not verified yet
        login_staff(self.client)
        self.client.post(reverse("payment_set", args=[self.reg_id()]), {"verified": "on"})
        self.assertEqual(len(self.certificate_mails()), 1)
        self.client.post(reverse("payment_set", args=[self.reg_id()]), {})  # untick, then tick again
        self.client.post(reverse("payment_set", args=[self.reg_id()]), {"verified": "on"})
        self.assertEqual(len(self.certificate_mails()), 1)  # never sent twice automatically

    def test_no_certificate_before_workshop_starts_or_without_template_or_if_rejected(self):
        self.post(); login_staff(self.client); rid = self.reg_id()
        self.client.post(reverse("payment_set", args=[rid]), {"verified": "on"})
        self.fb(email="Test@Example.com")  # workshop has not started and there is no template
        self.assertEqual(self.certificate_mails(), [])
        self.add_template(); self.start_workshop()
        self.client.post(reverse("payment_set", args=[rid]), {"status": "rejected"})
        self.client.post(reverse("certificate_send", args=[rid]))  # not complete: refused without "force"
        self.assertEqual(self.certificate_mails(), [])
        self.client.post(reverse("payment_set", args=[rid]), {"verified": "on"})
        self.assertEqual(len(self.certificate_mails()), 1)

    def test_free_workshop_certificate_needs_only_feedback(self):
        self.w.fee = None; self.w.save(); self.add_template()
        self.post(payment_screenshot=""); self.start_workshop()
        self.fb(email="Test@Example.com")
        self.assertEqual(len(self.certificate_mails()), 1)

    def test_staff_can_resend_and_override(self):
        self.add_template(); self.post(); self.start_workshop(); login_staff(self.client); rid = self.reg_id()
        self.client.post(reverse("certificate_send", args=[rid], ), {"force": "1"})  # no payment or feedback, staff override
        self.assertEqual(len(self.certificate_mails()), 1)
        self.client.post(reverse("certificate_send", args=[rid]), {"force": "1"})  # resend
        self.assertEqual(len(self.certificate_mails()), 2)
        pdf = self.client.get(reverse("certificate_download", args=[rid]))
        self.assertEqual(pdf["Content-Type"], "application/pdf")

    def test_certificate_email_failure_is_recorded_and_retryable(self):
        self.add_template(); self.post(); self.start_workshop(); login_staff(self.client); rid = self.reg_id()
        self.client.post(reverse("payment_set", args=[rid]), {"verified": "on"})
        with mock.patch("workshops.mailer.send_certificate", side_effect=RuntimeError("smtp down")):
            self.fb(email="Test@Example.com")
        self.assertTrue(str(self.cert_status()).startswith("failed"))
        self.client.post(reverse("certificate_send", args=[rid]))
        self.assertEqual(self.cert_status(), "sent")

    def test_certificate_setup_preview_and_permissions(self):
        url = reverse("certificate_setup", args=[self.w.pk])
        self.assertEqual(self.client.get(url).status_code, 302)
        self.assertEqual(self.client.get(reverse("certificate_preview", args=[self.w.pk])).status_code, 302)
        login_staff(self.client)
        self.assertEqual(self.client.get(reverse("certificate_preview", args=[self.w.pk])).status_code, 404)  # no template yet
        r = self.client.post(url, {"cert_template": SimpleUploadedFile("t.png", cert_template_bytes(), content_type="image/png"),
                                   "cert_name_x": "50", "cert_name_y": "45", "cert_font_size": "5", "cert_color": "#112233", "cert_font": ""})
        self.assertEqual(r.status_code, 302)
        self.w.refresh_from_db()
        self.assertTrue(self.w.has_certificate)
        self.assertEqual((self.w.cert_name_y, self.w.cert_color), (45.0, "#112233"))
        p = self.client.get(reverse("certificate_preview", args=[self.w.pk]) + "?name=Asha+Devi&x=40&y=55&size=6&color=%23ff0000")
        self.assertEqual(p["Content-Type"], "image/png"); self.assertTrue(p.content.startswith(b"\x89PNG"))
        self.assertEqual(self.client.get(reverse("certificate_preview", args=[self.w.pk]) + "?x=abc&size=nan&color=bad").status_code, 200)
        bad = self.client.post(url, {"cert_template": SimpleUploadedFile("t.png", b"nope", content_type="image/png"),
                                     "cert_name_x": "50", "cert_name_y": "45", "cert_font_size": "5", "cert_color": "#112233"})
        self.assertEqual(bad.status_code, 200)  # not an image: form shown again with an error
        self.assertNotIn("/media/", self.w.cert_template.name)

    def test_payment_checkbox_redirect_is_safe(self):
        self.post(); login_staff(self.client)
        r = self.client.post(reverse("payment_set", args=[self.reg_id()]), {"verified": "on", "next": "https://evil.example/x"})
        self.assertEqual(r["Location"], reverse("registrations"))
        r = self.client.post(reverse("payment_set", args=[self.reg_id()]), {"verified": "on", "next": "/dashboard/registrations/?q=test"})
        self.assertEqual(r["Location"], "/dashboard/registrations/?q=test")

    # --- announcing a new workshop to past students
    def past_workshop_with_students(self):
        now = timezone.now()
        past = Workshop.objects.create(title="Git Basics", start_datetime=now - timedelta(days=9), end_datetime=now - timedelta(days=9) + timedelta(hours=2), status="published")
        def reg(i, email, status="verified", wid=None):
            storage.add_registration({"registration_id": f"E2E-20260101-AAAAA{i}", "registered_at": f"2026-01-0{i} 10:00:00",
                                      "workshop_id": wid or past.pk, "workshop_title": "Git Basics", "full_name": f"Student {i}",
                                      "email": email, "registration_status": "confirmed", "payment_status": status}, None)
        reg(1, "old1@x.com"); reg(2, "old2@x.com"); reg(3, "gone@x.com"); reg(4, "rejected@x.com", status="rejected")
        reg(5, "already@x.com"); reg(5, "already@x.com", wid=self.w.pk)  # already signed up for the new workshop
        Unsubscribed.objects.create(email="gone@x.com")
        return past

    def test_announcement_goes_once_to_the_right_people(self):
        self.past_workshop_with_students(); login_staff(self.client)
        page = self.client.get(reverse("announce", args=[self.w.pk]))
        self.assertContains(page, "Past students")
        self.client.post(reverse("announce", args=[self.w.pk]), {"action": "send"})
        got = sorted(m.to[0] for m in mail.outbox if m.subject.startswith("New workshop"))
        self.assertEqual(got, ["old1@x.com", "old2@x.com"])
        body = [m for m in mail.outbox if m.to == ["old1@x.com"]][0]
        self.assertIn("/register/?w=" + self.w.slug, body.body)
        self.assertIn("/unsubscribe/?t=", body.body)
        self.assertNotIn("meet.example", body.body)  # meeting link is never in announcements
        self.assertEqual(Announcement.objects.filter(status="sent").count(), 2)
        n = len(mail.outbox)
        self.client.post(reverse("announce", args=[self.w.pk]), {"action": "send"})
        self.assertEqual(len(mail.outbox), n)  # nobody gets it twice
        self.w.refresh_from_db(); self.assertIsNotNone(self.w.announced_at)

    def test_announcement_failures_are_retried_and_drafts_refused(self):
        self.past_workshop_with_students(); login_staff(self.client)
        with mock.patch("workshops.mailer.send_announcement", side_effect=RuntimeError("smtp down")):
            self.client.post(reverse("announce", args=[self.w.pk]), {"action": "send"})
        self.assertEqual(Announcement.objects.filter(status__startswith="failed").count(), 2)
        self.client.post(reverse("announce", args=[self.w.pk]), {"action": "send"})
        self.assertEqual(Announcement.objects.filter(status="sent").count(), 2)
        self.w.status = "draft"; self.w.save()
        n = len(mail.outbox)
        self.client.post(reverse("announce", args=[self.w.pk]), {"action": "send"})
        self.assertEqual(len(mail.outbox), n)

    def test_unsubscribe_link_works_and_forged_links_do_not(self):
        from .announce import unsubscribe_url
        link = unsubscribe_url("Someone@X.com")
        self.assertContains(self.client.get(link), "someone@x.com")
        self.assertEqual(Unsubscribed.objects.count(), 0)  # a GET alone changes nothing
        token = link.split("t=")[1]
        self.assertContains(self.client.post(reverse("unsubscribe"), {"t": token}), "You are unsubscribed")
        self.assertTrue(Unsubscribed.objects.filter(email="someone@x.com").exists())
        self.assertContains(self.client.get(reverse("unsubscribe") + "?t=forged"), "not valid")

    def test_announce_requires_staff(self):
        self.assertEqual(self.client.post(reverse("announce", args=[self.w.pk]), {"action": "send"}).status_code, 302)
        self.assertEqual(len(mail.outbox), 0)

    # --- workshop list and gallery pages
    def test_workshops_page_lists_workshops_and_tells_staff_about_drafts(self):
        Workshop.objects.create(title="Hidden Draft", start_datetime=timezone.now() + timedelta(days=20), status="draft")
        page = self.client.get(reverse("workshop_list"))
        self.assertContains(page, "AI Prompt Engineering"); self.assertNotContains(page, "Hidden Draft")
        login_staff(self.client)
        self.assertContains(self.client.get(reverse("workshop_list")), "Hidden Draft")

    def test_home_with_only_the_featured_workshop_has_no_misleading_empty_box(self):
        home = self.client.get(reverse("home")).content.decode()
        self.assertNotIn("No other upcoming workshops", home)
        self.assertIn("See all workshops", home)

    def test_gallery_page_aligned_grid_markup(self):
        GalleryImage.objects.create(kind="attendees", image=png("a.png"), caption="Batch photo")
        GalleryImage.objects.create(kind="feedback", image=png("b.png"), caption="Chat feedback")
        page = self.client.get(reverse("gallery")).content.decode()
        self.assertIn("Batch photo", page); self.assertIn("gallery-tall", page)

    # --- certificate Yes/No, amounts, certificate IDs
    def two_students(self):
        self.w.capacity = None; self.w.save()
        self.post(full_name="First Student", email="one@example.com")
        self.post(full_name="Second Student", email="two@example.com")

    def col(self, row, name):
        ws = self.rows()
        headers = [c.value for c in ws[1]]
        return ws.cell(row=row, column=headers.index(name) + 1).value

    def test_certificate_choice_yes_pays_no_does_not(self):
        self.w.capacity = None; self.w.save()
        self.assertContains(self.post(wants_certificate=""), "Please choose Yes or No.")
        self.assertRedirects(self.post(wants_certificate="no", payment_screenshot="", email="no@example.com"), reverse("thank_you"))
        self.assertEqual((self.col(2, "wants_certificate"), self.col(2, "payment_status"), self.col(2, "payment_screenshot"), self.col(2, "amount")),
                         ("no", "none", None, None))
        self.assertRedirects(self.post(email="yes@example.com"), reverse("thank_you"))
        self.assertEqual((self.col(3, "wants_certificate"), self.col(3, "payment_status"), str(self.col(3, "amount"))), ("yes", "pending", "9"))
        self.assertContains(self.post(email="y2@example.com", payment_screenshot=""), "Attach your payment screenshot")

    def test_fee_everyone_must_pay_has_no_choice(self):
        self.w.fee_for_certificate = False; self.w.save()
        self.assertContains(self.post(wants_certificate="", payment_screenshot=""), "Attach your payment screenshot")
        self.assertRedirects(self.post(wants_certificate=""), reverse("thank_you"))
        self.assertEqual(self.col(2, "wants_certificate"), "yes")

    def test_register_page_offers_the_choice(self):
        self.assertContains(self.client.get(reverse("register")), '"choice": true')
        self.w.fee_for_certificate = False; self.w.save()
        self.assertContains(self.client.get(reverse("register")), '"choice": false')

    def test_no_certificate_is_ever_sent_to_someone_who_did_not_ask(self):
        self.add_template(); self.post(wants_certificate="no", payment_screenshot=""); self.start_workshop()
        self.fb(email="Test@Example.com")
        login_staff(self.client)
        self.client.post(reverse("certificate_send", args=[self.reg_id()]))
        self.assertEqual(self.certificate_mails(), [])
        self.assertNotContains(self.client.get(reverse("registrations")), "Send now")

    def test_certificate_ids_count_up_and_stay_with_the_student(self):
        self.two_students(); self.add_template(); self.start_workshop(); login_staff(self.client)
        r = self.client.post(reverse("certificate_setup", args=[self.w.pk]), {
            "cert_name_x": "50", "cert_name_y": "50", "cert_font_size": "5", "cert_color": "#112233", "cert_font": "",
            "cert_id_x": "50", "cert_id_y": "90", "cert_id_size": "1.6", "cert_id_start": "E2E/WS/2026/500"})
        self.assertEqual(r.status_code, 302)
        self.w.refresh_from_db()
        self.assertEqual((self.w.cert_id_prefix, self.w.cert_id_next), ("E2E/WS/2026/", 500))
        ids = [self.col(2, "registration_id"), self.col(3, "registration_id")]
        for i, em in enumerate(("one@example.com", "two@example.com")):
            self.client.post(reverse("payment_set", args=[ids[i]]), {"verified": "on"})
            self.fb(email=em)
        self.assertEqual([self.col(2, "certificate_id"), self.col(3, "certificate_id")], ["E2E/WS/2026/500", "E2E/WS/2026/501"])
        self.assertEqual(len(self.certificate_mails()), 2)
        self.assertIn("E2E/WS/2026/500", self.certificate_mails()[0].body)
        self.client.post(reverse("certificate_send", args=[ids[0]]), {"force": "1"})  # resend keeps the same number
        self.assertEqual(self.col(2, "certificate_id"), "E2E/WS/2026/500")
        self.w.refresh_from_db(); self.assertEqual(self.w.cert_id_next, 502)
        page = self.client.get(reverse("certificate_setup", args=[self.w.pk]))
        self.assertContains(page, "E2E/WS/2026/502")
        r = self.client.post(reverse("certificate_setup", args=[self.w.pk]), {
            "cert_name_x": "50", "cert_name_y": "50", "cert_font_size": "5", "cert_color": "#112233", "cert_font": "",
            "cert_id_x": "50", "cert_id_y": "90", "cert_id_size": "1.6", "cert_id_start": "E2E/WS/2026/502"})
        self.w.refresh_from_db(); self.assertEqual(self.w.cert_id_next, 502)  # saving without editing does not restart
        bad = self.client.post(reverse("certificate_setup", args=[self.w.pk]), {
            "cert_name_x": "50", "cert_name_y": "50", "cert_font_size": "5", "cert_color": "#112233", "cert_font": "",
            "cert_id_x": "50", "cert_id_y": "90", "cert_id_size": "1.6", "cert_id_start": "no number here"})
        self.assertContains(bad, "must end with a number")

    def test_payments_page_adds_up_the_amounts(self):
        self.two_students(); login_staff(self.client)
        self.client.post(reverse("payment_set", args=[self.col(2, "registration_id")]), {"verified": "on"})
        page = self.client.get(reverse("payments"))
        self.assertContains(page, "Verified"); self.assertContains(page, "AI Prompt Engineering")
        t = page.context["totals"]
        self.assertEqual((int(t["verified"]), int(t["pending"]), t["screenshots"]), (9, 9, 2))
        self.assertContains(self.client.get(reverse("dashboard")), "Amount verified")
        self.client.logout()
        self.assertEqual(self.client.get(reverse("payments")).status_code, 302)

    def test_students_filter_by_confirmation_email_sent_or_not(self):
        self.w.capacity = None; self.w.save()
        self.post(full_name="Mailed Student", email="m@example.com")
        with mock.patch("workshops.mailer.send_all", return_value="failed: smtp down"):
            self.post(full_name="Unmailed Student", email="u@example.com")
        login_staff(self.client)
        bad = self.client.get(reverse("registrations") + "?email=notsent").content.decode()
        self.assertIn("Unmailed Student", bad); self.assertNotIn("Mailed Student</div>", bad.replace("Unmailed Student", ""))
        good = self.client.get(reverse("registrations") + "?email=sent").content.decode()
        self.assertIn("Mailed Student", good); self.assertNotIn("Unmailed Student", good)
        self.post(full_name="No Cert", email="n@example.com", wants_certificate="no", payment_screenshot="")
        only_no = self.client.get(reverse("registrations") + "?cert=no").content.decode()
        self.assertIn("No Cert", only_no); self.assertNotIn("Mailed Student", only_no)

    def test_staff_can_upload_a_payment_screenshot_by_hand(self):
        self.post(wants_certificate="no", payment_screenshot=""); login_staff(self.client); rid = self.reg_id()
        bad = SimpleUploadedFile("x.png", b"not an image", content_type="image/png")
        self.client.post(reverse("payment_upload", args=[rid]), {"screenshot": bad})
        self.assertEqual(self.col(2, "payment_status"), "none")
        self.client.post(reverse("payment_upload", args=[rid]), {"screenshot": png(), "verified": "on"})
        self.assertEqual((self.col(2, "payment_status"), self.col(2, "wants_certificate")), ("verified", "yes"))
        name = self.col(2, "payment_screenshot")
        self.assertTrue(name and (TMP / "payments" / name).is_file())
        self.assertEqual(self.client.get(reverse("payment_proof", args=[rid])).status_code, 200)
        self.client.logout()
        self.assertEqual(self.client.post(reverse("payment_upload", args=[rid]), {"screenshot": png()}).status_code, 302)

    def test_qr_maker_makes_a_qr_for_any_link(self):
        self.assertEqual(self.client.get(reverse("qr_admin")).status_code, 302)
        login_staff(self.client)
        page = self.client.get(reverse("qr_admin") + "?link=https%3A%2F%2Fexample.com%2Fform&caption=Join")
        self.assertContains(page, "custom.png")
        img = self.client.get(reverse("qr_custom_png") + "?u=https://example.com/form&caption=Join")
        self.assertTrue(img.content.startswith(b"\x89PNG"))
        self.assertContains(self.client.get(reverse("qr_admin") + "?link=javascript%3Aalert(1)"), "http://")
        self.assertEqual(self.client.get(reverse("qr_custom_png") + "?u=javascript:alert(1)").status_code, 404)

    def test_announcement_uses_my_text_and_brochures(self):
        self.past_workshop_with_students(); login_staff(self.client)
        pdf = SimpleUploadedFile("brochure.pdf", b"%PDF-1.4 sample", content_type="application/pdf")
        exe = SimpleUploadedFile("virus.exe", b"MZ....", content_type="application/octet-stream")
        fake = SimpleUploadedFile("fake.pdf", b"this is not a pdf", content_type="application/pdf")
        self.client.post(reverse("announce", args=[self.w.pk]), {
            "action": "save", "subject": "Join our AI workshop", "message": "Hello {name}, come and learn!",
            "files": [pdf, png("flyer.png"), exe, fake]})
        draft = self.w.announce_draft
        self.assertEqual(sorted(f.original_name for f in draft.files.all()), ["brochure.pdf", "flyer.png"])  # bad files refused
        self.assertEqual(len(mail.outbox), 0)  # saving sends nothing
        self.client.post(reverse("announce", args=[self.w.pk]), {"action": "send", "subject": "Join our AI workshop", "message": "Hello {name}, come and learn!"})
        got = [m for m in mail.outbox if m.subject == "Join our AI workshop"]
        self.assertEqual(sorted(m.to[0] for m in got), ["old1@x.com", "old2@x.com"])
        one = [m for m in got if m.to == ["old1@x.com"]][0]
        self.assertIn("Hello Student 1, come and learn!", one.body)
        self.assertEqual(sorted(a[0] for a in one.attachments), ["brochure.pdf", "flyer.png"])
        gone = draft.files.first()
        self.client.post(reverse("announce", args=[self.w.pk]), {"action": "save", "subject": "S", "message": "M", "remove": [str(gone.pk)]})
        self.assertEqual(draft.files.count(), 1)
        self.assertContains(self.client.get(reverse("announce", args=[self.w.pk])), "already registered")

    def test_organisers_are_emailed_when_the_last_seat_goes(self):
        self.w.capacity = 2; self.w.save()
        self.post(email="a@example.com")
        self.assertEqual([m for m in mail.outbox if m.subject.startswith("Seats full")], [])
        self.post(email="b@example.com")
        alerts = [m for m in mail.outbox if m.subject.startswith("Seats full")]
        self.assertEqual(len(alerts), 1); self.assertEqual(alerts[0].to, ["admin@example.com"])

    def test_confirmation_email_goes_to_the_student(self):
        self.post(wants_certificate="no", payment_screenshot="")
        to_student = [m for m in mail.outbox if m.to == ["test@example.com"]]
        self.assertEqual(len(to_student), 1)
        self.assertIn("without a certificate", to_student[0].alternatives[0][0])

    def test_countdown_blinking_seats_and_poster_frames(self):
        self.w.poster.save("p.png", ContentFile(png().read()), save=True)
        for url in (reverse("home"), self.w.get_absolute_url(), reverse("workshop_list")):
            html = self.client.get(url).content.decode()
            self.assertIn("poster-frame", html, url)
        for url in (reverse("home"), self.w.get_absolute_url()):
            html = self.client.get(url).content.decode()
            self.assertIn('class="countdown"', html); self.assertIn("data-deadline", html)
            self.assertIn("seats-low", html)  # 1 seat left blinks
        self.w.capacity = 50; self.w.save()
        self.assertNotIn("seats-low", self.client.get(reverse("home")).content.decode())
        Workshop.objects.filter(pk=self.w.pk).update(capacity=1, start_datetime=timezone.now() - timedelta(hours=1), end_datetime=timezone.now() + timedelta(hours=1))
        self.assertNotIn('class="countdown"', self.client.get(self.w.get_absolute_url()).content.decode())  # closed: no timer

    # --- model rules
    def test_only_one_featured(self):
        w2 = Workshop.objects.create(title="Two", start_datetime=timezone.now() + timedelta(days=9), status="published", featured=True)
        self.w.refresh_from_db()
        self.assertFalse(self.w.featured); self.assertTrue(w2.featured)

    def test_unique_slugs(self):
        a = Workshop.objects.create(title="Same", start_datetime=self.w.start_datetime)
        b = Workshop.objects.create(title="Same", start_datetime=self.w.start_datetime)
        self.assertNotEqual(a.slug, b.slug)


def tearDownModule():
    shutil.rmtree(TMP, ignore_errors=True)
