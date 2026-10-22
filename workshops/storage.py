"""
Registrations are stored ONLY in one Excel file (settings.REGISTRATIONS_XLSX).

Safety rules implemented here:
  * Every read/modify/write happens while holding a file lock, so two students submitting at the same
    moment cannot overwrite each other. Duplicate and seat-limit checks run inside that same lock.
  * Saves are atomic: the workbook is written to a temporary file, then swapped in. A crash cannot leave
    a half-written registrations file.
  * Cell values are always stored as text, so nothing a student types can run as an Excel formula.
  * If the file is open in Excel on Windows, the swap is refused and the student sees an honest error
    instead of a false "success".
"""
import io
import os
from contextlib import contextmanager
from pathlib import Path

from django.conf import settings
from filelock import FileLock, Timeout
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

SHEET = "Registrations"
HEADERS = [
    "registration_id", "registered_at", "workshop_id", "workshop_title", "full_name", "email", "phone",
    "college", "department", "year_of_study", "participation_mode", "consent", "registration_status",
    "email_status", "source", "payment_reference", "payment_screenshot",
    # added later; always appended on the right so older files keep working
    "payment_status", "certificate_status", "certificate_sent_at",
    "current_address", "education_qualification", "current_status", "heard_from", "interest_domain",
    "skill_level", "demo_interest", "demo_mode", "demo_slot", "queries", "referred_by",
    "wants_certificate", "amount", "certificate_id",
]
WIDTHS = [22, 20, 12, 34, 24, 30, 14, 34, 18, 20, 18, 9, 18, 34, 10, 20, 30,
          16, 30, 20, 34, 24, 22, 26, 26, 18, 30, 14, 14, 36, 24, 18, 10, 24]
assert len(HEADERS) == len(WIDTHS)


class StorageError(Exception):
    """The Excel file could not be read or written. The message is safe to show to users."""


class DuplicateError(StorageError):
    pass


class FullError(StorageError):
    pass


def path() -> Path:
    return Path(settings.REGISTRATIONS_XLSX)


@contextmanager
def _locked():
    p = path()
    p.parent.mkdir(parents=True, exist_ok=True)
    lock = FileLock(str(p) + ".lock", timeout=20)
    try:
        with lock:
            yield
    except Timeout:
        raise StorageError("The registration file is busy. Please try again in a moment.")


def _style(ws):
    for i, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1B8A3F")
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[c.column_letter].width = WIDTHS[i - 1]
    ws.freeze_panes = "A2"


def _atomic_save(wb, target: Path):
    tmp = target.with_name(target.stem + ".saving" + target.suffix)
    try:
        wb.save(tmp)
        os.replace(tmp, target)
    except PermissionError:
        _silent_remove(tmp)
        raise StorageError(
            "The registrations file is open in Excel or another program. "
            "Close it and try again (use the dashboard's Download button instead of opening the live file)."
        )
    except OSError:
        _silent_remove(tmp)
        raise StorageError("The registration could not be saved to the Excel file. Please try again.")


def _silent_remove(p: Path):
    try:
        p.unlink()
    except OSError:
        pass


def _open():
    p = path()
    if not p.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET
        ws.append(HEADERS)
        _style(ws)
        _atomic_save(wb, p)
    try:
        wb = load_workbook(p)
    except Exception:
        raise StorageError("The registrations file could not be read. It may be damaged or still being written.")
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.active
    headers = [c.value for c in ws[1]]
    for h in HEADERS:  # add any missing columns on the right; never reorder or delete existing ones
        if h not in headers:
            headers.append(h)
            ws.cell(row=1, column=len(headers), value=h)
    return wb, ws, headers


def _rows(ws, headers):
    out = []
    for vals in ws.iter_rows(min_row=2, values_only=True):
        row = {h: ("" if v is None else v) for h, v in zip(headers, vals)}
        if row.get("registration_id"):
            out.append(row)
    return out


def _put(ws, cells_values):
    """Append a row, forcing every text cell to stay text (no formula execution)."""
    ws.append(cells_values)
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"
            cell.number_format = "@"


def add_registration(row: dict, capacity=None):
    """Save one registration. Raises DuplicateError / FullError / StorageError. Returns after a durable save."""
    wid = str(row["workshop_id"])
    email = str(row["email"]).strip().lower()
    with _locked():
        wb, ws, headers = _open()
        same = [r for r in _rows(ws, headers) if str(r["workshop_id"]) == wid and r["registration_status"] == "confirmed"]
        if any(str(r["email"]).strip().lower() == email for r in same):
            raise DuplicateError("Already registered")
        if capacity and len(same) >= int(capacity):
            raise FullError("Workshop full")
        _put(ws, [row.get(h, "") for h in headers])
        _atomic_save(wb, path())


def list_registrations():
    with _locked():
        wb, ws, headers = _open()
        return _rows(ws, headers)


def count_confirmed():
    counts = {}
    for r in list_registrations():
        if r["registration_status"] == "confirmed":
            k = str(r["workshop_id"])
            counts[k] = counts.get(k, 0) + 1
    return counts


def update_fields(registration_id, **changes):
    """Change columns of one registration. Returns True if the registration was found."""
    with _locked():
        wb, ws, headers = _open()
        r = _find_row(ws, headers, registration_id)
        if not r:
            return False
        for name, value in changes.items():
            if name in headers:
                c = ws.cell(row=r, column=headers.index(name) + 1)
                c.value = value
                c.data_type = "s"
        _atomic_save(wb, path())
        return True


def update_email_status(registration_id, status):
    return update_fields(registration_id, email_status=status)


def _find_row(ws, headers, registration_id):
    id_col = headers.index("registration_id") + 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=id_col).value == registration_id:
            return r
    return None


def get_registration(registration_id):
    for row in list_registrations():
        if row["registration_id"] == registration_id:
            return row
    return None


def payment_state(row) -> str:
    """verified / pending / rejected / free / none (no certificate wanted). Older rows are worked out."""
    st = str(row.get("payment_status", "") or "").strip()
    if st:
        return st
    return "pending" if row.get("payment_screenshot") else "free"


SENDING_TIMEOUT_MIN = 10


def claim_certificate(registration_id, force=False):
    """Atomically mark a certificate as 'sending' so two clicks or two triggers never send it twice.

    Returns the registration row if this caller may send it, or None if it was already sent or is being sent.
    """
    from datetime import datetime, timedelta
    with _locked():
        wb, ws, headers = _open()
        r = _find_row(ws, headers, registration_id)
        if not r:
            return None
        row = {h: ("" if ws.cell(row=r, column=i + 1).value is None else ws.cell(row=r, column=i + 1).value)
               for i, h in enumerate(headers)}
        st = str(row.get("certificate_status", "")).strip()
        if st.startswith("sent") and not force:
            return None
        if st.startswith("sending"):
            try:
                started = datetime.fromisoformat(st.split(" ", 1)[1])
                if datetime.now() - started < timedelta(minutes=SENDING_TIMEOUT_MIN):
                    return None
            except (IndexError, ValueError):
                pass  # unreadable marker: treat as stale
        c = ws.cell(row=r, column=headers.index("certificate_status") + 1)
        c.value = "sending " + datetime.now().isoformat(timespec="seconds")
        c.data_type = "s"
        _atomic_save(wb, path())
        return row


def finish_certificate(registration_id, status):
    from datetime import datetime
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S") if status.startswith("sent") else ""
    return update_fields(registration_id, certificate_status=status, certificate_sent_at=stamp)


def rows_to_xlsx(rows) -> bytes:
    """Build a fresh workbook (for downloads) from already-filtered rows."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(HEADERS)
    _style(ws)
    for r in rows:
        _put(ws, [r.get(h, "") for h in HEADERS])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
